#!/usr/bin/env python3
"""Reproduce the inter-rater agreement statistics reported in the review.

The script aligns the two reviewers by study and year, excludes Fan et al.
(2024) from the final 29-study dataset, treats the Figure 5 axes as ordinal
1--5 ratings, and excludes Figure 6 NA cells pairwise. Confidence intervals
are obtained by resampling studies (not individual cells) with replacement.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
N_BOOTSTRAP = 10_000
RANDOM_SEED = 2026


def normalize(value):
    if value is None:
        return None
    text = str(value).strip()
    if text.upper() in {"NA", "N/A", "NOT APPLICABLE", ""}:
        return None
    return int(float(text))


def read_ratings(filename: str, columns: list[str]):
    workbook = load_workbook(ROOT / filename, data_only=True, read_only=True)
    sheet = workbook["Independent ratings"]
    headers = {str(cell.value).strip(): index for index, cell in enumerate(sheet[1], 1)}
    records = {}
    for row in range(2, sheet.max_row + 1):
        study = sheet.cell(row, headers["Study"]).value
        year = sheet.cell(row, headers["Year"]).value
        if not study:
            continue
        if str(study).strip().lower().startswith("fan et al") and str(year).strip() == "2024":
            continue
        key = (str(study).strip(), str(year).strip())
        records[key] = [normalize(sheet.cell(row, headers[column]).value) for column in columns]
    return records


def paired_arrays(left, right, keys, column_indices):
    a, b = [], []
    for key in keys:
        for column in column_indices:
            x, y = left[key][column], right[key][column]
            if x is not None and y is not None:
                a.append(x)
                b.append(y)
    return np.asarray(a, dtype=int), np.asarray(b, dtype=int)


def weighted_kappa(a, b, categories, weighting="quadratic"):
    categories = np.asarray(categories, dtype=int)
    index = {value: i for i, value in enumerate(categories)}
    observed = np.zeros((len(categories), len(categories)), dtype=float)
    for x, y in zip(a, b):
        observed[index[int(x)], index[int(y)]] += 1
    observed /= observed.sum()
    row = observed.sum(axis=1)
    col = observed.sum(axis=0)
    expected = np.outer(row, col)
    distance = np.abs(np.subtract.outer(np.arange(len(categories)), np.arange(len(categories))))
    if weighting == "unweighted":
        weights = (distance > 0).astype(float)
    elif weighting == "linear":
        weights = distance / (len(categories) - 1)
    elif weighting == "quadratic":
        weights = (distance / (len(categories) - 1)) ** 2
    else:
        raise ValueError(f"Unknown weighting: {weighting}")
    denominator = np.sum(weights * expected)
    return 1.0 - np.sum(weights * observed) / denominator


def cluster_bootstrap(left, right, keys, column_indices, categories, weighting):
    rng = np.random.default_rng(RANDOM_SEED)
    estimates = []
    for _ in range(N_BOOTSTRAP):
        sampled = [keys[i] for i in rng.integers(0, len(keys), len(keys))]
        a, b = paired_arrays(left, right, sampled, column_indices)
        estimates.append(weighted_kappa(a, b, categories, weighting))
    return np.percentile(estimates, [2.5, 97.5])


def report(label, left, right, columns, column_indices, categories):
    keys = sorted(set(left) & set(right))
    a, b = paired_arrays(left, right, keys, column_indices)
    print(f"\n{label}: {len(keys)} studies; {len(a)} paired ratings")
    print(f"Exact agreement: {(a == b).sum()}/{len(a)}")
    for weighting in ("unweighted", "linear", "quadratic"):
        kappa = weighted_kappa(a, b, categories, weighting)
        low, high = cluster_bootstrap(
            left, right, keys, column_indices, categories, weighting
        )
        print(f"{weighting:11s} kappa = {kappa:.2f} (95% CI {low:.2f}-{high:.2f})")


def main():
    fig5_columns = ["Validation", "Readiness"]
    fig5_a = read_ratings("Azza_Figure5_scoring.xlsx", fig5_columns)
    fig5_p = read_ratings("Pooja_Figure5_scoring.xlsx", fig5_columns)
    report("Figure 5: validation depth", fig5_a, fig5_p, fig5_columns, [0], range(1, 6))
    report("Figure 5: deployment readiness", fig5_a, fig5_p, fig5_columns, [1], range(1, 6))
    report("Figure 5: pooled axes", fig5_a, fig5_p, fig5_columns, [0, 1], range(1, 6))

    fig6_columns = ["F1", "F2", "E1", "E2", "E3", "E4", "R1", "V1", "V2"]
    fig6_a = read_ratings("Azza_Figure6_scoring.xlsx", fig6_columns)
    fig6_p = read_ratings("Pooja_Figure6_scoring.xlsx", fig6_columns)
    report("Figure 6: nine reporting items", fig6_a, fig6_p, fig6_columns, list(range(9)), range(3))


if __name__ == "__main__":
    main()
